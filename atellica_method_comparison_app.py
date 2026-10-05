import io
import re
import copy
import fitz
import openpyxl
import streamlit as st

st.set_page_config(page_title="Atellica Method Comparison", layout="wide")

RESULT_RE = re.compile(
    r"^([<>]?\s*[0-9]+(?:\.[0-9]+)?)\s*"
    r"(g/dL|mg/dL|ug/dL|µg/dL|ng/dL|U/L|mEq/L|mmol/L|µmol/L)?$",
    re.I
)
SID_RE = re.compile(r"(?i)^b2b[0-9]+(?:\.\.)?$")


def parse_result(line):
    m = RESULT_RE.match(line.strip())
    if not m:
        return None, None
    number = m.group(1).replace(" ", "")
    unit = m.group(2) or ""
    raw = f"{number} {unit}".strip()
    if number.startswith("<") or number.startswith(">"):
        return raw, None
    try:
        return raw, float(number)
    except ValueError:
        return raw, None


def extract_pdf(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    lines = []
    for page_no, page in enumerate(doc, 1):
        for raw in page.get_text().splitlines():
            line = raw.strip()
            if line:
                lines.append((page_no, line))

    headings = []
    for i, (_, line) in enumerate(lines):
        if line == "Patient Name" and i > 0:
            heading = lines[i - 1][1]
            if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{1,30}", heading):
                if heading not in headings:
                    headings.append(heading)

    data = {}
    for i, (_, line) in enumerate(lines):
        if line != "Patient Name" or i == 0:
            continue

        assay = lines[i - 1][1]
        if assay not in headings:
            continue

        j = i + 1
        while j < len(lines):
            if (
                lines[j][1] == "Patient Name"
                and j > 0
                and lines[j - 1][1] in headings
            ):
                break
            j += 1

        block = lines[i:j]
        rows = []
        last_result_index = -1

        for k, (_, text) in enumerate(block):
            raw, value = parse_result(text)
            if raw is None:
                continue

            sid = ""
            for p in range(k - 1, max(-1, last_result_index), -1):
                candidate = block[p][1]
                if SID_RE.match(candidate):
                    sid = candidate
                    break

            rows.append({
                "sid": sid,
                "key": re.sub(r"\s+", "", sid).upper(),
                "raw": raw,
                "value": value,
            })
            last_result_index = k

        if rows:
            data.setdefault(assay, []).extend(rows)

    return data


def normalize_assay_name(name):
    return re.sub(r"\s+", "", name).upper()


def common_parameters(old_data, new_data):
    old_norm = {}
    new_norm = {}

    for name in old_data:
        old_norm[normalize_assay_name(name)] = name

    for name in new_data:
        new_norm[normalize_assay_name(name)] = name

    common = []
    for norm, old_name in old_norm.items():
        if norm in new_norm:
            common.append((old_name, new_norm[norm]))
    return common


def safe_sheet_name(name, existing):
    base = re.sub(r'[:\\/?*\[\]]', "_", name or "Parameter")[:31]
    if not base:
        base = "Parameter"

    candidate = base
    counter = 2
    while candidate in existing:
        suffix = f"_{counter}"
        candidate = base[:31 - len(suffix)] + suffix
        counter += 1
    return candidate


def copy_template_sheet(template_bytes, workbook, sheet_name):
    source_wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    source = source_wb[source_wb.sheetnames[0]]
    ws = workbook.create_sheet(sheet_name)

    for row in source.iter_rows():
        for source_cell in row:
            target = ws[source_cell.coordinate]
            target.value = source_cell.value

            # FIX: StyleArray does not provide .copy().
            # Use Python's copy.copy() for OpenPyXL style objects.
            if source_cell.has_style:
                target._style = copy.copy(source_cell._style)

            target.number_format = source_cell.number_format

            if source_cell.alignment:
                target.alignment = copy.copy(source_cell.alignment)

            if source_cell.protection:
                target.protection = copy.copy(source_cell.protection)

    for key, dimension in source.column_dimensions.items():
        ws.column_dimensions[key].width = dimension.width
        ws.column_dimensions[key].hidden = dimension.hidden

    for key, dimension in source.row_dimensions.items():
        ws.row_dimensions[key].height = dimension.height
        ws.row_dimensions[key].hidden = dimension.hidden

    for merged_range in source.merged_cells.ranges:
        ws.merge_cells(str(merged_range))

    return ws


def populate_workbook(
    template_bytes,
    old_data,
    new_data,
    common_pairs,
    metadata,
):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    existing_sheets = set()

    for old_assay, new_assay in common_pairs:
        parameter_name = old_assay
        sheet_name = safe_sheet_name(parameter_name, existing_sheets)
        existing_sheets.add(sheet_name)

        ws = copy_template_sheet(template_bytes, workbook, sheet_name)

        ws["B1"] = parameter_name
        ws["B2"] = metadata["instrument"]
        ws["B3"] = metadata["serial"]
        ws["B4"] = metadata["date"]

        ws["A5"] = metadata["reagent"]
        ws["B5"] = metadata["reagent_lot"]
        ws["C5"] = metadata["reagent_expiry"]

        ws["D5"] = metadata["calibrator"]
        ws["E5"] = metadata["calibrator_lot"]
        ws["F5"] = metadata["calibrator_expiry"]

        ws["G5"] = metadata["qc_lot"]
        ws["H5"] = metadata["qc_expiry"]

        ws["C8"] = old_assay
        ws["D8"] = new_assay
        ws["E8"] = "Y-X"
        ws["F8"] = "%(Y-X)/X"

        ws["C9"] = old_assay
        ws["D9"] = new_assay
        ws["E9"] = "Bias"
        ws["F9"] = "%Bias"

        ws["G9"] = "Mean"
        ws["H9"] = '=IFERROR(AVERAGE(C10:C29),"")'
        ws["I9"] = '=IFERROR(AVERAGE(D10:D29),"")'
        ws["J9"] = '=IFERROR((I9-H9)/H9*100,"")'

        ws["G10"] = "Slope"
        ws["H10"] = '=IFERROR(SLOPE(D10:D29,C10:C29),"")'

        ws["G11"] = "Intercept"
        ws["H11"] = '=IFERROR(INTERCEPT(D10:D29,C10:C29),"")'

        ws["G12"] = "R square"
        ws["H12"] = '=IFERROR(RSQ(D10:D29,C10:C29),"")'

        old_rows = old_data.get(old_assay, [])
        new_rows = new_data.get(new_assay, [])

        old_lookup = {row["key"]: row for row in old_rows if row["key"]}
        new_lookup = {row["key"]: row for row in new_rows if row["key"]}

        ordered_ids = []
        seen = set()

        for row in old_rows:
            if row["key"] and row["key"] not in seen:
                ordered_ids.append(row["key"])
                seen.add(row["key"])

        for row in new_rows:
            if row["key"] and row["key"] not in seen:
                ordered_ids.append(row["key"])
                seen.add(row["key"])

        for excel_row in range(10, 30):
            ws.cell(excel_row, 1).value = excel_row - 9
            ws.cell(excel_row, 2).value = None
            ws.cell(excel_row, 3).value = None
            ws.cell(excel_row, 4).value = None

            index = excel_row - 10
            if index < len(ordered_ids):
                sample_key = ordered_ids[index]
                old_row = old_lookup.get(sample_key)
                new_row = new_lookup.get(sample_key)

                if old_row:
                    ws.cell(excel_row, 2).value = old_row["sid"]
                    ws.cell(excel_row, 3).value = (
                        old_row["value"]
                        if old_row["value"] is not None
                        else old_row["raw"]
                    )

                if new_row:
                    if not old_row:
                        ws.cell(excel_row, 2).value = new_row["sid"]
                    ws.cell(excel_row, 4).value = (
                        new_row["value"]
                        if new_row["value"] is not None
                        else new_row["raw"]
                    )

            ws.cell(
                excel_row, 5
            ).value = (
                f'=IF(AND(ISNUMBER(C{excel_row}),'
                f'ISNUMBER(D{excel_row})),'
                f'D{excel_row}-C{excel_row},"")'
            )

            ws.cell(
                excel_row, 6
            ).value = (
                f'=IF(AND(ISNUMBER(C{excel_row}),'
                f'ISNUMBER(D{excel_row}),'
                f'C{excel_row}<>0),'
                f'(D{excel_row}-C{excel_row})/C{excel_row}*100,"")'
            )

        ws.freeze_panes = "A10"

    return workbook


st.title("Atellica 2 vs Atellica 1 — Method Comparison")

st.info(
    "Two PDFs are mandatory. OLD PDF = Atellica 2. "
    "NEW PDF = Atellica 1. "
    "The SAME parameter name in both PDFs is compared with itself. "
    "Example: Alb vs Alb and ALB_A vs ALB_A."
)

old_file = st.file_uploader(
    "1. OLD Instrument PDF — Atellica 2",
    type=["pdf"],
)

new_file = st.file_uploader(
    "2. NEW Instrument PDF — Atellica 1",
    type=["pdf"],
)

template_file = st.file_uploader(
    "3. Your Method Comparison Excel Template",
    type=["xlsx"],
)

if old_file and new_file and template_file:
    with st.spinner("Extracting all parameters from both PDFs..."):
        old_data = extract_pdf(old_file.getvalue())
        new_data = extract_pdf(new_file.getvalue())
        common_pairs = common_parameters(old_data, new_data)

    st.subheader("Parameters to be compared")

    comparison_table = []
    for old_assay, new_assay in common_pairs:
        comparison_table.append({
            "Atellica 2 PDF": old_assay,
            "Atellica 1 PDF": new_assay,
            "Old values": len(old_data.get(old_assay, [])),
            "New values": len(new_data.get(new_assay, [])),
            "Comparison": f"{old_assay}  vs  {new_assay}",
        })

    st.dataframe(
        comparison_table,
        use_container_width=True,
        hide_index=True,
    )

    st.success(
        f"{len(common_pairs)} SAME-NAME parameters found in BOTH PDFs."
    )

    st.caption(
        "No _2 → _A conversion is used. "
        "If the parameter is Alb, the comparison is Alb vs Alb. "
        "If the parameter is ALB_A, the comparison is ALB_A vs ALB_A."
    )

    st.subheader("Excel template information")

    c1, c2, c3 = st.columns(3)

    instrument = c1.text_input(
        "Instrument",
        "Atellica 2 / Atellica 1",
    )
    serial = c2.text_input("Serial No.", "")
    report_date = c3.text_input("Date", "09/30/2026")

    c1, c2, c3 = st.columns(3)
    reagent = c1.text_input("Reagent Name", "")
    reagent_lot = c2.text_input("Reagent Lot", "")
    reagent_expiry = c3.text_input("Reagent Expiry", "")

    c1, c2, c3 = st.columns(3)
    calibrator = c1.text_input("Calibrator Name", "")
    calibrator_lot = c2.text_input("Calibrator Lot", "")
    calibrator_expiry = c3.text_input("Calibrator Expiry", "")

    c1, c2 = st.columns(2)
    qc_lot = c1.text_input("QC Lot", "")
    qc_expiry = c2.text_input("QC Expiry", "")

    if st.button(
        "POPULATE ATTACHED EXCEL TEMPLATE",
        type="primary",
        use_container_width=True,
    ):
        metadata = {
            "instrument": instrument,
            "serial": serial,
            "date": report_date,
            "reagent": reagent,
            "reagent_lot": reagent_lot,
            "reagent_expiry": reagent_expiry,
            "calibrator": calibrator,
            "calibrator_lot": calibrator_lot,
            "calibrator_expiry": calibrator_expiry,
            "qc_lot": qc_lot,
            "qc_expiry": qc_expiry,
        }

        with st.spinner(
            "Populating the exact Excel template for every common parameter..."
        ):
            workbook = populate_workbook(
                template_file.getvalue(),
                old_data,
                new_data,
                common_pairs,
                metadata,
            )
            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)

        st.success(
            f"Completed. {len(common_pairs)} parameters populated."
        )

        st.download_button(
            "DOWNLOAD POPULATED METHOD COMPARISON EXCEL",
            data=output.getvalue(),
            file_name="Atellica_2_vs_Atellica_1_Method_Comparison.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )
